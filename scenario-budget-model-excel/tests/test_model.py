"""Tests that need no Excel: the Python mirror's invariants and the workbook's structure.

(The Excel-vs-Python numerical comparison needs desktop Excel; see src/finalize.ps1 and src/verify.py.)
"""
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
import build_workbook as bw  # noqa: E402
import model_spec as ms  # noqa: E402


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    out = tmp_path_factory.mktemp("xl") / "model.xlsx"
    bw.OUT = out
    bw.build()
    return openpyxl.load_workbook(out)


@pytest.mark.parametrize("scenario", ms.SCENARIOS)
def test_cash_and_customer_roll_forwards_hold(scenario):
    r = ms.engine(ms.drivers(scenario), ms.HIRES[scenario])
    assert r["close_cash"][-1] == pytest.approx(ms.drivers(scenario)["open_cash"] + sum(r["net_cash"]))
    assert r["close_cust"][-1] == ms.drivers(scenario)["open_customers"] + sum(r["new"]) - sum(r["churned"])
    for i in range(12):
        assert r["ebitda"][i] == pytest.approx(r["revenue"][i] - r["cogs"][i] - r["opex"][i])
        assert r["net_income"][i] == pytest.approx(r["ebit"][i] - r["tax"][i])


def test_scenarios_are_ordered_sensibly():
    rev = {s: sum(ms.engine(ms.drivers(s), ms.HIRES[s])["revenue"]) for s in ms.SCENARIOS}
    assert rev["Downside"] < rev["Base"] < rev["Upside"]
    cash = {s: ms.engine(ms.drivers(s), ms.HIRES[s])["close_cash"][-1] for s in ms.SCENARIOS}
    assert cash["Downside"] < cash["Base"] < cash["Upside"]


def test_round_half_up_matches_excel_rounding():
    assert ms.round_half_up(2.5) == 3 and ms.round_half_up(0.5) == 1 and ms.round_half_up(2.4999) == 2


def test_actuals_are_deterministic_and_complete():
    a1, a2 = ms.synthetic_actuals(), ms.synthetic_actuals()
    assert a1 == a2
    assert all(len(v) == ms.N_ACTUAL for v in a1.values())


def test_workbook_sheets_names_and_validation(built):
    assert built.sheetnames == ["Cover", "Dashboard", "Assumptions", "Budget_vs_Actual", "Actuals", "Budget_Locked", "Model_Base",
                                "Model_Upside", "Model_Downside", "Scenarios", "ChartData", "Checks"]
    assert {"Scenario", "Scn_Idx", "Check_Status"} <= set(built.defined_names.keys())
    dv = built["Dashboard"].data_validations.dataValidation
    assert any("Base,Upside,Downside" in (v.formula1 or "") for v in dv)


def test_engine_sheets_have_no_hardcoded_numbers_in_month_cells(built):
    """Every month cell on the three live engine sheets must be a formula (except the structural month-number row)."""
    for name in ("Model_Base", "Model_Upside", "Model_Downside"):
        ws = built[name]
        for key, row in bw.ROW.items():
            if key == "idx":
                continue
            for col in bw.MCOLS:
                v = ws[f"{col}{row}"].value
                assert isinstance(v, str) and v.startswith("="), f"{name}!{col}{row} ({key}) is not a formula: {v!r}"


def test_locked_budget_is_values_not_formulas(built):
    ws = built["Budget_Locked"]
    for key in ("revenue", "ebitda", "close_cash"):
        for col in bw.MCOLS:
            assert not isinstance(ws[f"{col}{bw.ROW[key]}"].value, str)


def test_locked_budget_equals_base_case(built):
    base = ms.engine(ms.drivers("Base"), ms.HIRES["Base"])
    ws = built["Budget_Locked"]
    for i, col in enumerate(bw.MCOLS):
        assert ws[f"{col}{bw.ROW['revenue']}"].value == pytest.approx(base["revenue"][i])
